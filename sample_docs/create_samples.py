import docx
from pypdf import PdfWriter

# Create sample DOCX file
doc = docx.Document()
doc.add_heading('Cloud Security Compliance Policy 2026', 0)

doc.add_heading('Section 1: Data Encryption Standards', level=1)
doc.add_paragraph('All sensitive user data stored at rest must utilize AES-256 encryption. Encryption keys must be rotated every 90 days following ISO 27001 security standards.')

doc.add_heading('Section 2: Access Control & IAM', level=1)
doc.add_paragraph('Multi-factor authentication (MFA) is strictly required for all administrative access. Least privilege access controls must be audited quarterly.')

doc.save('sample_docs/Cloud_Security_Policy.docx')
print('Sample DOCX created successfully at sample_docs/Cloud_Security_Policy.docx')
