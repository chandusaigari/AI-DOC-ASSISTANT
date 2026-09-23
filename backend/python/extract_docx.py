import sys
import json
import os
import docx

def extract_docx(file_path):
    if not os.path.exists(file_path):
        return {"success": False, "error": f"File not found: {file_path}"}
    
    try:
        doc = docx.Document(file_path)
        sections_data = []
        full_text_list = []
        
        current_heading = "General"
        para_counter = 1
        
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
                
            if p.style and p.style.name.startswith("Heading"):
                current_heading = text
            
            sections_data.append({
                "section": current_heading,
                "para_index": para_counter,
                "text": text
            })
            para_counter += 1
            full_text_list.append(text)
            
        # Also extract table texts if present
        for t_idx, table in enumerate(doc.tables):
            table_texts = []
            for row in table.rows:
                row_str = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_str:
                    table_texts.append(row_str)
            if table_texts:
                t_content = "\n".join(table_texts)
                sections_data.append({
                    "section": f"Table {t_idx + 1}",
                    "para_index": para_counter,
                    "text": t_content
                })
                para_counter += 1
                full_text_list.append(t_content)

        full_text = "\n\n".join(full_text_list)
        return {
            "success": True,
            "filename": os.path.basename(file_path),
            "total_sections": len(sections_data),
            "sections": sections_data,
            "full_text": full_text
        }
    except Exception as e:
        return {"success": False, "error": str(e)}

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(json.dumps({"success": False, "error": "No file path provided"}))
        sys.exit(1)
        
    path = sys.argv[1]
    result = extract_docx(path)
    print(json.dumps(result, ensure_ascii=False))
